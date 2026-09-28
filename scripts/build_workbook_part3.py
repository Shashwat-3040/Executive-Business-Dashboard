import calendar
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

PATH = "/home/claude/executive-dashboard/excel/Executive_Business_Performance_Dashboard.xlsx"

FONT_NAME = "Arial"
NAVY = "1F3864"
WHITE = "FFFFFF"
header_font = Font(name=FONT_NAME, bold=True, color=WHITE, size=10)
header_fill = PatternFill("solid", fgColor=NAVY)

def style_header_row(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

def autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

wb = load_workbook(PATH)
regions = ["Central", "East", "North", "South", "West"]
departments = ["Electronics", "Apparel", "Home & Living", "Grocery", "Beauty & Personal Care"]

ws = wb.create_sheet("Chart_Data")

# --- Monthly revenue trend (24 rows) ---
ws.cell(row=1, column=1, value="MonthLabel")
ws.cell(row=1, column=2, value="Year")
ws.cell(row=1, column=3, value="Month")
ws.cell(row=1, column=4, value="NetRevenue")
style_header_row(ws, 1, 4)
row = 2
for y in (2024, 2025):
    for m in range(1, 13):
        ws.cell(row=row, column=1, value=f"{calendar.month_abbr[m]}-{str(y)[2:]}")
        ws.cell(row=row, column=2, value=y)
        ws.cell(row=row, column=3, value=m)
        ws.cell(row=row, column=4,
                 value=f"=SUMIFS('Sales_Summary'!$D$2:$D$121,'Sales_Summary'!$A$2:$A$121,B{row},'Sales_Summary'!$B$2:$B$121,C{row})")
        ws.cell(row=row, column=4).number_format = "$#,##0"
        row += 1
trend_last = row - 1

# --- Region gross margin % vs company avg (5 rows), via INDEX/MATCH against Margin_Leakage ---
r0 = row + 1
ws.cell(row=r0, column=1, value="Region")
ws.cell(row=r0, column=2, value="GrossMarginPct")
ws.cell(row=r0, column=3, value="CompanyAvgMarginPct")
style_header_row(ws, r0, 3)
row = r0 + 1
region_margin_start = row
for reg in regions:
    ws.cell(row=row, column=1, value=reg)
    ws.cell(row=row, column=2,
             value=f"=INDEX('Margin_Leakage'!$D$2:$D$6,MATCH(A{row},'Margin_Leakage'!$A$2:$A$6,0))")
    ws.cell(row=row, column=3,
             value=f"=INDEX('Margin_Leakage'!$E$2:$E$6,MATCH(A{row},'Margin_Leakage'!$A$2:$A$6,0))")
    ws.cell(row=row, column=2).number_format = "0.0%"
    ws.cell(row=row, column=3).number_format = "0.0%"
    row += 1
region_margin_end = row - 1

# --- Revenue by department (5 rows), direct SUMIFS from Raw_Sales ---
r0 = row + 1
ws.cell(row=r0, column=1, value="Department")
ws.cell(row=r0, column=2, value="NetRevenue")
style_header_row(ws, r0, 2)
row = r0 + 1
dept_rev_start = row
for dept in departments:
    ws.cell(row=row, column=1, value=dept)
    ws.cell(row=row, column=2,
             value=f"=SUMIFS('Raw_Sales'!$O$2:$O$110700,'Raw_Sales'!$G$2:$G$110700,A{row})")
    ws.cell(row=row, column=2).number_format = "$#,##0"
    row += 1
dept_rev_end = row - 1

# --- Churn rate by region (5 rows), INDEX/MATCH against the Churn "ALL TIERS" rollup rows (17-21) ---
r0 = row + 1
ws.cell(row=r0, column=1, value="Region")
ws.cell(row=r0, column=2, value="ChurnRatePct")
style_header_row(ws, r0, 2)
row = r0 + 1
churn_reg_start = row
for reg in regions:
    ws.cell(row=row, column=1, value=reg)
    ws.cell(row=row, column=2,
             value=f"=INDEX('Churn'!$E$17:$E$21,MATCH(A{row},'Churn'!$A$17:$A$21,0))")
    ws.cell(row=row, column=2).number_format = "0.0%"
    row += 1
churn_reg_end = row - 1

# --- Inventory turnover by department (5 rows), INDEX/MATCH against "ALL REGIONS" rollup rows (27-31) ---
r0 = row + 1
ws.cell(row=r0, column=1, value="Department")
ws.cell(row=r0, column=2, value="TurnoverRatio")
style_header_row(ws, r0, 2)
row = r0 + 1
turn_dept_start = row
for dept in departments:
    ws.cell(row=row, column=1, value=dept)
    ws.cell(row=row, column=2,
             value=f"=INDEX('Inventory_Turnover'!$E$27:$E$31,MATCH(A{row},'Inventory_Turnover'!$B$27:$B$31,0))")
    ws.cell(row=row, column=2).number_format = "0.00"
    row += 1
turn_dept_end = row - 1

autosize(ws, [16, 8, 8, 14])
wb.save(PATH)
print("Chart_Data written. Ranges:",
      dict(trend_last=trend_last, region_margin=(region_margin_start, region_margin_end),
           dept_rev=(dept_rev_start, dept_rev_end), churn_reg=(churn_reg_start, churn_reg_end),
           turn_dept=(turn_dept_start, turn_dept_end)))
