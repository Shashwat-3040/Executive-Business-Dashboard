from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.utils import get_column_letter

PATH = "/home/claude/executive-dashboard/excel/Executive_Business_Performance_Dashboard.xlsx"

FONT_NAME = "Arial"
NAVY = "1F3864"
BLUE = "2E5395"
LIGHT_BLUE = "D9E2F3"
GOLD = "BF8F00"
RED = "C00000"
WHITE = "FFFFFF"

title_font = Font(name=FONT_NAME, bold=True, size=18, color=NAVY)
subtitle_font = Font(name=FONT_NAME, italic=True, size=11, color="595959")
label_font = Font(name=FONT_NAME, bold=True, size=10)
normal_font = Font(name=FONT_NAME, size=10)
kpi_value_font = Font(name=FONT_NAME, bold=True, size=20, color=NAVY)
kpi_value_font_red = Font(name=FONT_NAME, bold=True, size=20, color=RED)
kpi_label_font = Font(name=FONT_NAME, size=10, color="595959")
card_fill = PatternFill("solid", fgColor=LIGHT_BLUE)
alert_fill = PatternFill("solid", fgColor="FCE4E4")
input_fill = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="BFBFBF")
box_border = Border(left=thin, right=thin, top=thin, bottom=thin)

def autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

wb = load_workbook(PATH)
ws = wb.create_sheet("Dashboard", 0)  # make it the first/active sheet
ws.sheet_view.showGridLines = False

ws.merge_cells("B2:N2")
ws["B2"] = "Executive Business Performance 360° Dashboard"
ws["B2"].font = title_font
ws.merge_cells("B3:N3")
ws["B3"] = "NorthStar Retail Group  |  FY2024 – FY2025  |  5 Regions · 40 Stores · 300 SKUs · 6,000 Loyalty Customers"
ws["B3"].font = subtitle_font

# ---------------------------------------------------------------- KPI cards
kpi_defs = [
    ("Total Net Revenue\n(FY24-FY25)", "=EBITDA!C12", "$#,##0,,\"M\"", False),
    ("YoY Revenue Growth", "=(SUMIFS(EBITDA!C2:C11,EBITDA!B2:B11,2025)-SUMIFS(EBITDA!C2:C11,EBITDA!B2:B11,2024))/SUMIFS(EBITDA!C2:C11,EBITDA!B2:B11,2024)", "0.0%", False),
    ("Company EBITDA Margin", "=EBITDA!J12", "0.0%", False),
    ("Overall Customer Churn", "=Churn!E22", "0.0%", False),
    ("Avg. Inventory Turnover", "=AVERAGE(Inventory_Turnover!E2:E26)", "0.00\"x/yr\"", False),
    ("South Region Margin Leakage", "=Margin_Leakage!G5", "0.0%", True),
]
col = 2
row_top = 5
row_bot = 8
for label, formula, fmt, alert in kpi_defs:
    c1 = get_column_letter(col)
    c2 = get_column_letter(col + 1)
    fill = alert_fill if alert else card_fill
    # label row (merged)
    ws.merge_cells(f"{c1}{row_top}:{c2}{row_top+1}")
    lbl_cell = ws[f"{c1}{row_top}"]
    lbl_cell.value = label
    lbl_cell.font = kpi_label_font
    lbl_cell.fill = fill
    lbl_cell.border = box_border
    lbl_cell.alignment = Alignment(wrap_text=True, vertical="bottom")
    # value row (merged)
    ws.merge_cells(f"{c1}{row_top+2}:{c2}{row_bot}")
    val_cell = ws[f"{c1}{row_top+2}"]
    val_cell.value = formula
    val_cell.font = kpi_value_font_red if alert else kpi_value_font
    val_cell.number_format = fmt
    val_cell.fill = fill
    val_cell.border = box_border
    val_cell.alignment = Alignment(vertical="center")
    col += 2
autosize(ws, [3] + [13] * 13)
for r in (5, 6, 7, 8):
    ws.row_dimensions[r].height = 18

# ---------------------------------------------------------------- Customer lookup tool
lr = 11
ws.cell(row=lr, column=2, value="Customer Lookup").font = label_font
ws.cell(row=lr + 1, column=2, value="Enter Customer ID:").font = normal_font
lookup_cell = f"C{lr+1}"
ws[lookup_cell] = 42
ws[lookup_cell].fill = input_fill
ws[lookup_cell].border = box_border

fields = [
    ("Name", f"=IFERROR(INDEX('Dim_Customer'!$B$2:$B$6001,MATCH({lookup_cell},'Dim_Customer'!$A$2:$A$6001,0)),\"Not found\")"),
    ("Region", f"=IFERROR(INDEX('Dim_Customer'!$C$2:$C$6001,MATCH({lookup_cell},'Dim_Customer'!$A$2:$A$6001,0)),\"\")"),
    ("Membership Tier", f"=IFERROR(INDEX('Dim_Customer'!$D$2:$D$6001,MATCH({lookup_cell},'Dim_Customer'!$A$2:$A$6001,0)),\"\")"),
    ("Join Date", f"=IFERROR(INDEX('Dim_Customer'!$E$2:$E$6001,MATCH({lookup_cell},'Dim_Customer'!$A$2:$A$6001,0)),\"\")"),
    ("Status", f"=IFERROR(IF(INDEX('Dim_Customer'!$F$2:$F$6001,MATCH({lookup_cell},'Dim_Customer'!$A$2:$A$6001,0))=\"\",\"Active\",\"Churned\"),\"\")"),
]
for i, (flabel, formula) in enumerate(fields):
    ws.cell(row=lr + 2 + i, column=2, value=flabel).font = normal_font
    ws.cell(row=lr + 2 + i, column=3, value=formula).font = normal_font
ws.cell(row=lr, column=2).alignment = Alignment(horizontal="left")
ws.cell(row=lr - 1, column=2, value="Built with INDEX/MATCH (functionally interchangeable with XLOOKUP)").font = subtitle_font

# ================================================================= Charts
chart_row_start = 20

# 1. Monthly net revenue trend (line)
lc = LineChart()
lc.title = "Monthly Net Revenue Trend"
lc.style = 2
lc.y_axis.title = "Net Revenue ($)"
lc.x_axis.title = "Month"
lc.height = 8
lc.width = 18
data = Reference(wb["Chart_Data"], min_col=4, min_row=1, max_row=25)
cats = Reference(wb["Chart_Data"], min_col=1, min_row=2, max_row=25)
lc.add_data(data, titles_from_data=True)
lc.set_categories(cats)
lc.series[0].smooth = False
ws.add_chart(lc, f"B{chart_row_start}")

# 2. Gross margin % by region (bar) vs company avg (line) — combo chart
bc = BarChart()
bc.title = "Gross Margin % by Region (South = margin leakage)"
bc.style = 10
bc.y_axis.title = "Gross Margin %"
bc.height = 8
bc.width = 18
bar_data = Reference(wb["Chart_Data"], min_col=2, min_row=27, max_row=32)
line_data = Reference(wb["Chart_Data"], min_col=3, min_row=27, max_row=32)
bcats = Reference(wb["Chart_Data"], min_col=1, min_row=28, max_row=32)
bc.add_data(bar_data, titles_from_data=True)
bc.set_categories(bcats)
lc2 = LineChart()
lc2.add_data(line_data, titles_from_data=True)
lc2.series[0].smooth = False
bc += lc2
ws.add_chart(bc, f"H{chart_row_start}")

# 3. Revenue by department (pie)
pc = PieChart()
pc.title = "Revenue Mix by Department"
pc.height = 8
pc.width = 12
pdata = Reference(wb["Chart_Data"], min_col=2, min_row=34, max_row=39)
pcats = Reference(wb["Chart_Data"], min_col=1, min_row=35, max_row=39)
pc.add_data(pdata, titles_from_data=True)
pc.set_categories(pcats)
pc.dataLabels = DataLabelList()
pc.dataLabels.showPercent = True
ws.add_chart(pc, f"B{chart_row_start + 17}")

# 4. Churn rate by region (bar)
bc2 = BarChart()
bc2.title = "Churn Rate % by Region"
bc2.style = 12
bc2.y_axis.title = "Churn Rate %"
bc2.height = 8
bc2.width = 12
cdata = Reference(wb["Chart_Data"], min_col=2, min_row=41, max_row=46)
ccats = Reference(wb["Chart_Data"], min_col=1, min_row=42, max_row=46)
bc2.add_data(cdata, titles_from_data=True)
bc2.set_categories(ccats)
ws.add_chart(bc2, f"G{chart_row_start + 17}")

# 5. Inventory turnover by department (bar)
bc3 = BarChart()
bc3.title = "Annual Inventory Turnover by Department"
bc3.style = 11
bc3.y_axis.title = "Turnover Ratio (x/yr)"
bc3.height = 8
bc3.width = 12
tdata = Reference(wb["Chart_Data"], min_col=2, min_row=48, max_row=53)
tcats = Reference(wb["Chart_Data"], min_col=1, min_row=49, max_row=53)
bc3.add_data(tdata, titles_from_data=True)
bc3.set_categories(tcats)
ws.add_chart(bc3, f"L{chart_row_start + 17}")

wb.save(PATH)
print("Dashboard sheet with charts written.")
